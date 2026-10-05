#!/usr/bin/env python3
"""
scripts/sync_pins_to_firebase.py
Đồng bộ mã PIN và Token xác thực sang Firebase Realtime Database.
Đảm bảo mã PIN tuyệt đối không bị lộ ra tệp docs/data.json công khai.
"""
import os
import sys
import json
import hashlib
from typing import Dict, Any

_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _root not in sys.path:
    sys.path.insert(0, _root)

if sys.platform.startswith("win"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

from src.sync_firebase import push_tokens_to_firebase

def hash_str(val: str) -> str:
    return hashlib.sha256(str(val).strip().encode("utf-8")).hexdigest()

def main():
    print("=" * 65)
    print("🔥 CHUYỂN ĐỔI & ĐỒNG BỘ MÃ PIN SANG FIREBASE REALTIME DATABASE")
    print("=" * 65)

    tokens: Dict[str, Any] = {}

    # 1. Thêm Admin & Class mật khẩu mặc định
    admin_hash = hash_str("THAYPHUNG2026")
    tokens[admin_hash] = {"role": "admin", "name": "Quản trị viên / Giáo viên", "class": "ALL"}
    tokens[hash_str("CPP2026")] = {"role": "class", "name": "Lớp C++", "class": "C++"}
    tokens[hash_str("PYTHON2026")] = {"role": "class", "name": "Lớp Python", "class": "Python"}
    tokens[hash_str("VIP11")] = {"role": "class", "name": "Lớp Python 1-1", "class": "Python 1-1"}

    # 2. Đọc từ auth_tokens_for_firebase.json nếu đã có
    backup_file = os.path.join(_root, "auth_tokens_for_firebase.json")
    if os.path.exists(backup_file):
        try:
            with open(backup_file, "r", encoding="utf-8") as f:
                d = json.load(f)
                existing_tokens = d.get("auth_tokens", {})
                tokens.update(existing_tokens)
                print(f"📦 Đã nạp {len(existing_tokens)} token từ {backup_file}")
        except Exception as e:
            print(f"⚠️ Lỗi đọc backup: {e}")

    # 3. Đọc từ file Excel nếu có
    excel_path = os.path.join(_root, "Quản lý học sinh.xlsx")
    if os.path.exists(excel_path):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            for sname in ["Học Sinh", "Danh sách Học sinh", "Students"]:
                if sname in wb.sheetnames:
                    ws = wb[sname]
                    # Tìm cột pin
                    header = [str(ws.cell(1, c).value or "").lower() for c in range(1, ws.max_column + 1)]
                    pin_col = -1
                    name_col = 2
                    cls_col = 3
                    for idx, h in enumerate(header, 1):
                        if any(k in h for k in ["pin", "mật khẩu", "pass"]):
                            pin_col = idx
                        elif "họ" in h or "tên" in h:
                            name_col = idx
                        elif "lớp" in h:
                            cls_col = idx
                    
                    if pin_col != -1:
                        stu_count = 0
                        for r in range(2, ws.max_row + 1):
                            name_val = str(ws.cell(r, name_col).value or "").strip()
                            pin_val = str(ws.cell(r, pin_col).value or "").strip()
                            cls_val = str(ws.cell(r, cls_col).value or "C++").strip()
                            if name_val and pin_val:
                                h_pin = hash_str(pin_val)
                                tokens[h_pin] = {
                                    "role": "student",
                                    "name": name_val,
                                    "class": cls_val,
                                    "stt": r - 1
                                }
                                stu_count += 1
                        print(f"📗 Đã trích xuất {stu_count} mã PIN học sinh từ file Excel '{sname}'")
                    break
        except Exception as e:
            print(f"ℹ️ Không đọc được file Excel: {e}")

    print(f"\n🔐 Tổng cộng chuẩn bị đồng bộ {len(tokens)} tokens xác thực.")

    # 4. Ghi file JSON cho Firebase Import
    with open(backup_file, "w", encoding="utf-8") as f:
        json.dump({"auth_tokens": tokens}, f, ensure_ascii=False, indent=2)
    print(f"💾 Đã lưu tệp Import Firebase tại: {backup_file}")

    # 5. Đẩy lên Firebase nếu có cấu hình
    db_url = sys.argv[1] if len(sys.argv) > 1 else None
    success = push_tokens_to_firebase(tokens, db_url=db_url)

    if not success:
        print("\n" + "-" * 65)
        print("📋 HƯỚNG DẪN 1 PHÚT NHẬP DỮ LIỆU LÊN FIREBASE:")
        print("1. Truy cập: https://console.firebase.google.com/")
        print("2. Vào Realtime Database > Tab 'Data'")
        print("3. Nhấp vào biểu tượng dấu 3 chấm (⋮) ở góc trên bên phải")
        print("4. Chọn 'Import JSON' (Nhập JSON)")
        print(f"5. Chọn tệp: {backup_file}")
        print("6. Vào tab 'Rules', dán nội dung trong file firebase_rules.json")
        print("7. Mở file docs/firebase-config.js và dán URL Realtime Database của bạn!")
        print("-" * 65)

if __name__ == "__main__":
    main()
